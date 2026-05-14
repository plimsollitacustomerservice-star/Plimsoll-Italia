"""Template.xlsx inspection utilities.

The application treats Template.xlsx as the source of truth. These helpers read
labels and formula metadata dynamically; they do not hard-code row numbers.
"""

from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from italian_accounts_reclassifier.mapping.normalizer import normalize_italian_label
from italian_accounts_reclassifier.models import TemplateTarget

REQUIRED_SHEETS = ("Form", "Manual")


@dataclass(frozen=True)
class TemplateInspection:
    """Summary of discovered Template.xlsx writable targets and formulas."""

    template_path: Path
    template_hash: str
    targets: tuple[TemplateTarget, ...]
    form_label_count: int
    manual_label_count: int
    formula_count: int


def workbook_hash(path: str | Path) -> str:
    """Return a content hash for template version tracking."""
    digest = sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _label_from_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def inspect_template(template_path: str | Path) -> TemplateInspection:
    """Inspect Form and Manual target labels and formulas from a local workbook."""
    path = Path(template_path)
    workbook = load_workbook(path, data_only=False, read_only=False)
    missing = [sheet for sheet in REQUIRED_SHEETS if sheet not in workbook.sheetnames]
    if missing:
        raise ValueError(f"Template workbook is missing required sheet(s): {', '.join(missing)}")

    targets: list[TemplateTarget] = []
    formula_count = 0
    form_label_count = 0
    manual_label_count = 0

    form = workbook["Form"]
    formula_rows: set[int] = set()
    for row in range(1, form.max_row + 1):
        value = form.cell(row=row, column=4).value
        if isinstance(value, str) and value.startswith("="):
            formula_count += 1
            formula_rows.add(row)

    for sheet_name, writable_column in (("Form", "B"), ("Manual", "B")):
        sheet = workbook[sheet_name]
        for row in range(1, sheet.max_row + 1):
            label = _label_from_cell(sheet.cell(row=row, column=1).value)
            if not label:
                continue
            if sheet_name == "Form":
                form_label_count += 1
            else:
                manual_label_count += 1
            targets.append(
                TemplateTarget(
                    template_sheet=sheet_name,
                    template_row=row,
                    template_column=writable_column,
                    template_label=label,
                    normalized_template_label=normalize_italian_label(label),
                    destination_cell=f"{writable_column}{row}",
                    is_formula_related=sheet_name == "Form" and row in formula_rows,
                )
            )

    workbook.close()
    return TemplateInspection(
        template_path=path,
        template_hash=workbook_hash(path),
        targets=tuple(targets),
        form_label_count=form_label_count,
        manual_label_count=manual_label_count,
        formula_count=formula_count,
    )
