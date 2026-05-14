"""Template.xlsx copy-and-write export support."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook


@dataclass(frozen=True)
class ApprovedMappingValue:
    """A reviewed value ready to write into a copied Template.xlsx workbook."""

    sheet: str
    row: int
    value: float | int | str | None
    column: str = "B"

    @property
    def cell(self) -> str:
        return f"{self.column}{self.row}"


def write_template_copy(
    template_path: str | Path,
    output_path: str | Path,
    values: list[ApprovedMappingValue],
) -> Path:
    """Copy Template.xlsx, write approved values, and preserve workbook formulas."""
    source = Path(template_path)
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    copy2(source, destination)

    workbook = load_workbook(destination, data_only=False)
    formulas_before = _collect_formulas(workbook)
    for item in values:
        if item.sheet not in workbook.sheetnames:
            raise ValueError(f"Output workbook does not contain sheet: {item.sheet}")
        workbook[item.sheet][item.cell] = item.value

    formulas_after = _collect_formulas(workbook)
    if formulas_before != formulas_after:
        raise ValueError("Formula preservation check failed while writing Template.xlsx copy")

    workbook.save(destination)
    workbook.close()
    return destination


def _collect_formulas(workbook) -> dict[tuple[str, str], str]:
    formulas: dict[tuple[str, str], str] = {}
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formulas[(sheet.title, cell.coordinate)] = value
    return formulas
