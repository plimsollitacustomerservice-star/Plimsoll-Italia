"""Local XLSX extraction for account-label/value tables."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from italian_accounts_reclassifier.extractors.base import LocalExtractor
from italian_accounts_reclassifier.mapping.normalizer import normalize_italian_label
from italian_accounts_reclassifier.mapping.number_parser import parse_financial_number
from italian_accounts_reclassifier.models import ExtractedItem


class XlsxExtractor(LocalExtractor):
    """Extract simple adjacent label/value pairs from local .xlsx workbooks."""

    supported_extensions = (".xlsx",)

    def extract(self, path: Path) -> list[ExtractedItem]:
        workbook = load_workbook(path, data_only=True, read_only=True)
        items: list[ExtractedItem] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                values = [cell.value for cell in row]
                for index in range(len(values) - 1):
                    label = "" if values[index] is None else str(values[index]).strip()
                    parsed = parse_financial_number(values[index + 1])
                    if not label or parsed is None:
                        continue
                    source_column = get_column_letter(index + 2)
                    items.append(
                        ExtractedItem(
                            source_label=label,
                            normalized_label=normalize_italian_label(label),
                            value=str(values[index + 1]),
                            normalized_value=parsed,
                            detected_year=None,
                            source_type="xlsx",
                            extraction_method="xlsx_adjacent_label_value",
                            extraction_confidence=90.0,
                            source_sheet=sheet.title,
                            source_row=row[0].row,
                            source_column=source_column,
                            source_evidence={
                                "sheet": sheet.title,
                                "row": row[0].row,
                                "column": source_column,
                            },
                            raw_context_text=" | ".join("" if value is None else str(value) for value in values),
                        )
                    )
        workbook.close()
        return items
